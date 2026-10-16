"""Parse the 비교파트 Google Sheet export into a flat list of complexes (raw.json)."""
import json, re, sys, datetime
import openpyxl

SRC = sys.argv[1]
OUT = sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)

def clean(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m')
    s = str(v).strip()
    return s or None

AREAS = [ws.title for ws in wb.worksheets if ws.title not in ('키워드',)]

records = []
for ws in wb.worksheets:
    if ws.title == '키워드':
        continue
    area = ws.title
    if area == 'DMC가재울':
        # transposed: labels in column B, one complex per column from C
        labels = [clean(ws.cell(r, 2).value) for r in range(1, ws.max_row + 1)]
        name_row = 4  # row index (1-based) with names
        # find name row: first row where column C has a value and B empty
        for r in range(1, 10):
            if ws.cell(r, 3).value and not ws.cell(r, 2).value:
                name_row = r
                break
        for c in range(3, ws.max_column + 1):
            name = clean(ws.cell(name_row, c).value)
            if not name:
                continue
            rec = {'area': area, 'name': name, 'fields': {}}
            for r in range(name_row + 1, ws.max_row + 1):
                lab = labels[r - 1]
                if not lab:
                    continue
                v = ws.cell(r, c).value
                lab = lab.strip()
                # duplicated label 사우나 (community vs nearby) -> second is nearby
                key = lab
                if key in rec['fields']:
                    key = lab + '_2'
                rec['fields'][key] = clean(v)
            records.append(rec)
    else:
        hdr_rows = [r for r in range(1, 6) if clean(ws.cell(r, 1).value) == '아파트명']
        hr = hdr_rows[-1]
        labels = [clean(ws.cell(hr, c).value) for c in range(1, ws.max_column + 1)]
        for r in range(hr + 1, ws.max_row + 1):
            name = clean(ws.cell(r, 1).value)
            if not name or name == '아파트명':
                continue
            rec = {'area': area, 'name': name, 'fields': {}}
            for c in range(2, ws.max_column + 1):
                lab = labels[c - 1]
                if not lab:
                    continue
                lab = lab.strip()
                key = lab
                if key in rec['fields']:
                    key = lab + '_2'
                rec['fields'][key] = clean(ws.cell(r, c).value)
            records.append(rec)

# keyword sheet
kw = {}
ws = wb['키워드']
for row in ws.iter_rows(values_only=True):
    if not row[0] or not row[1]:
        continue
    kw[(clean(row[0]), clean(row[1]))] = {
        'score': row[2], 'level': clean(row[3]), 'keywords': clean(row[4]), 'link': clean(row[5])
    }
for rec in records:
    k = kw.get((rec['area'], rec['name']))
    rec['kw'] = k

json.dump(records, open(OUT, 'w'), ensure_ascii=False, indent=1)
print(len(records), 'records')
from collections import Counter
print(Counter(r['area'] for r in records))
print('missing kw', [(r['area'], r['name']) for r in records if not r['kw']][:40])
print('field keys', sorted({k for r in records for k in r['fields']}))
